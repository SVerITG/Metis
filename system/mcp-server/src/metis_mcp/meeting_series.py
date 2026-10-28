"""Meeting series — the continuity layer between one meeting and the next.

A meeting used to be a row that knew nothing about the meetings before it. The
`linked_meetings` column existed and nothing ever wrote it, so "follow up the
actions from last quarter" meant finding last quarter's report by hand and
re-reading it. This module makes the SERIES the thing that persists:

    meeting_series            one recurring meeting (e.g. a quarterly review with
                              one coordination). Holds the cadence, the language,
                              and the report rules learned from the researcher's
                              corrections — so a rule for one series never leaks
                              into another.
    meetings.series_id        each meeting belongs to at most one series, carries
                              a sequence label (T3 2026) and its predecessor.
    meeting_actions           one row per action, for its whole life. It gains an
                              owner and the series it belongs to.
    meeting_action_reviews    the status of an action AT a given meeting. One
                              action, many reviews — which is what a "status at
                              the next meeting" column is.

Recognising which series a new meeting continues is the hard part, because
recurring meetings with different groups look alike: the same template title, the
same agenda, and the researcher plus the same central colleagues in every one of
them. Title or topic similarity is therefore the WRONG primary signal. In order:

    1. the series code (a short organisation code) in the title or file names
    2. who attended, weighted so that someone present in every series counts for
       nothing and someone present in only one counts fully
    3. timing against the series' cadence
    4. title similarity, as a tie-breaker only

The matcher never links on its own. It proposes, and the caller asks the
researcher to confirm in one line. Linking a meeting to the wrong series attaches
another group's actions to its report, which is worse than no link at all.

Pure functions over a sqlite3 connection, with no MCP dependency, so the dashboard
and the MCP tools share one implementation instead of drifting apart.
"""
from __future__ import annotations

import datetime as _dt
import difflib
import json
import re
import sqlite3
import unicodedata
import uuid

# Kept in step with `system/installer/schema.sql`, which is the only mechanism
# that carries a schema change to the other computer on its own. `ensure_schema`
# exists because the MCP server can be the first process to touch a database the
# dashboard has not migrated yet.
_DDL_SERIES = """
CREATE TABLE IF NOT EXISTS meeting_series (
    series_id         TEXT PRIMARY KEY,
    name              TEXT NOT NULL,
    code              TEXT DEFAULT '',
    organisation      TEXT DEFAULT '',
    cadence_days      INTEGER DEFAULT 0,
    language          TEXT DEFAULT '',
    project           TEXT DEFAULT '',
    report_rules      TEXT DEFAULT '[]',
    notes             TEXT DEFAULT '',
    status            TEXT DEFAULT 'active',
    created_at        TEXT,
    updated_at        TEXT
)
"""

_DDL_REVIEWS = """
CREATE TABLE IF NOT EXISTS meeting_action_reviews (
    id          INTEGER PRIMARY KEY AUTOINCREMENT,
    action_id   INTEGER NOT NULL,
    meeting_id  TEXT NOT NULL,
    status      TEXT NOT NULL,
    note        TEXT DEFAULT '',
    reviewed_at TEXT
)
"""

_DDL_MEETING_ACTIONS = """
CREATE TABLE IF NOT EXISTS meeting_actions (
    id          INTEGER PRIMARY KEY AUTOINCREMENT,
    meeting_id  INTEGER NOT NULL,
    description TEXT NOT NULL,
    status      TEXT DEFAULT 'open',
    due_date    TEXT,
    created_at  TEXT
)
"""

_DDL_PERSONS = """
CREATE TABLE IF NOT EXISTS meeting_persons (
    person_id         TEXT PRIMARY KEY,
    name              TEXT,
    role              TEXT,
    last_meeting_date TEXT,
    meeting_count     INTEGER DEFAULT 0
)
"""

_DDL_ATTENDANCE = """
CREATE TABLE IF NOT EXISTS meeting_attendance (
    meeting_id TEXT,
    person_id  TEXT,
    PRIMARY KEY (meeting_id, person_id)
)
"""

# Columns added to tables that predate the series layer.
_ADDED_COLUMNS = {
    "meetings": [
        ("series_id", "TEXT DEFAULT ''"),
        ("sequence_label", "TEXT DEFAULT ''"),
        ("previous_meeting_id", "TEXT DEFAULT ''"),
        ("location", "TEXT DEFAULT ''"),
        ("summary", "TEXT DEFAULT ''"),
        ("discussion", "TEXT DEFAULT '[]'"),
        ("attachments", "TEXT DEFAULT '[]'"),
        ("report_path", "TEXT DEFAULT ''"),
        ("language", "TEXT DEFAULT ''"),
        ("next_meeting", "TEXT DEFAULT ''"),
        ("notes", "TEXT DEFAULT ''"),
        # Columns the dashboard adds at runtime; repeated here so a database the
        # MCP server touches first has them too.
        ("transcript", "TEXT"),
        ("duration_minutes", "INTEGER"),
        ("status", "TEXT DEFAULT 'filed'"),
    ],
    "meeting_actions": [
        ("owner", "TEXT DEFAULT ''"),
        ("series_id", "TEXT DEFAULT ''"),
        ("updated_at", "TEXT DEFAULT ''"),
        ("status_note", "TEXT DEFAULT ''"),
    ],
}

ACTION_STATUSES = ("open", "in_progress", "done", "dropped")
_CLOSED = ("done", "dropped")

# Thresholds for the matcher. A code match alone clears PROPOSE; attendance,
# timing and title together cannot, which is the point: without the code the
# researcher is always asked to choose.
_W_CODE_TITLE = 0.60
_W_CODE_TEXT = 0.25
_W_PEOPLE = 0.25
_W_TIMING = 0.10
_W_TITLE = 0.05
PROPOSE_AT = 0.45
CANDIDATE_AT = 0.20
MARGIN = 0.15


def _now() -> str:
    return _dt.datetime.now().isoformat(timespec="seconds")


def ensure_schema(conn: sqlite3.Connection) -> None:
    """Create the series tables and add missing columns. Idempotent."""
    for ddl in (_DDL_SERIES, _DDL_REVIEWS, _DDL_MEETING_ACTIONS, _DDL_PERSONS,
                _DDL_ATTENDANCE):
        conn.execute(ddl)
    has_meetings = conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name='meetings'"
    ).fetchone()
    if not has_meetings:
        conn.execute(
            "CREATE TABLE IF NOT EXISTS meetings (meeting_id TEXT PRIMARY KEY, "
            "title TEXT, meeting_date TEXT, domain TEXT, project TEXT, "
            "attendees TEXT, meeting_type TEXT, decisions TEXT, action_items TEXT, "
            "follow_ups TEXT, linked_meetings TEXT, transcript_status TEXT, "
            "transcript_path TEXT, created_at TEXT)"
        )
    for table, cols in _ADDED_COLUMNS.items():
        live = {r[1].lower() for r in conn.execute(f"PRAGMA table_info({table})")}
        for name, typedef in cols:
            if name not in live:
                conn.execute(f"ALTER TABLE {table} ADD COLUMN {name} {typedef}")
    conn.commit()


# ---------------------------------------------------------------------------
# Names and people
# ---------------------------------------------------------------------------

# Honorifics and roles that vary between transcripts of the same person.
_HONORIFICS = {
    "dr", "dre", "pr", "prof", "mr", "mrs", "ms", "mme", "mlle", "m", "mx",
    "monsieur", "madame", "docteur", "professeur", "sir",
}


def _fold(text: str) -> str:
    """Lowercase and strip accents, so 'Cécile' and 'Cecile' are one person."""
    text = unicodedata.normalize("NFKD", text or "")
    return "".join(c for c in text if not unicodedata.combining(c)).lower()


def person_key(name: str) -> str:
    """A stable identity for a participant name.

    Drops honorifics and anything in brackets ("Dr X (finance)") and sorts the
    remaining tokens, so word order in a transcript does not create a second
    person.
    """
    name = re.sub(r"\(.*?\)", " ", name or "")
    words = re.findall(r"[a-z0-9]+", _fold(name))
    tokens = [t for t in words if t not in _HONORIFICS and len(t) > 1]
    # A name made only of initials still names someone.
    return "-".join(sorted(tokens or [t for t in words if t not in _HONORIFICS]))


def parse_people(attendees) -> list[str]:
    """Participant names from a list, a JSON list, or a delimited string."""
    if not attendees:
        return []
    if isinstance(attendees, str):
        s = attendees.strip()
        if s.startswith("["):
            try:
                attendees = json.loads(s)
            except ValueError:
                attendees = re.split(r"[;,\n]", s)
        else:
            attendees = re.split(r"[;,\n]", s)
    out, seen = [], set()
    for a in attendees:
        if isinstance(a, dict):
            a = a.get("name") or ""
        a = re.sub(r"\s+", " ", str(a)).strip(" -•*\t")
        k = person_key(a)
        if a and k and k not in seen:
            seen.add(k)
            out.append(a)
    return out


# ---------------------------------------------------------------------------
# Series CRUD
# ---------------------------------------------------------------------------

def _rules(raw) -> list[str]:
    try:
        v = json.loads(raw or "[]")
        return [str(x) for x in v] if isinstance(v, list) else []
    except ValueError:
        return []


def create_series(conn, name: str, code: str = "", organisation: str = "",
                  cadence_days: int = 0, language: str = "", project: str = "",
                  report_rules: list[str] | None = None, notes: str = "") -> str:
    ensure_schema(conn)
    name = (name or "").strip()
    if not name:
        raise ValueError("A series needs a name.")
    code = (code or "").strip().upper()
    if code:
        clash = conn.execute(
            "SELECT series_id, name FROM meeting_series WHERE UPPER(code) = ?",
            (code,)).fetchone()
        if clash:
            raise ValueError(
                f"Code {code} already belongs to series '{clash['name']}' "
                f"({clash['series_id']}). A code must identify one series.")
    sid = f"ser-{uuid.uuid4().hex[:10]}"
    now = _now()
    conn.execute(
        "INSERT INTO meeting_series (series_id, name, code, organisation, "
        "cadence_days, language, project, report_rules, notes, status, "
        "created_at, updated_at) VALUES (?,?,?,?,?,?,?,?,?, 'active', ?, ?)",
        (sid, name, code, organisation.strip(), int(cadence_days or 0),
         language.strip(), project.strip(), json.dumps(report_rules or [],
                                                       ensure_ascii=False),
         notes.strip(), now, now))
    conn.commit()
    return sid


def get_series(conn, series_id: str) -> dict | None:
    ensure_schema(conn)
    row = conn.execute("SELECT * FROM meeting_series WHERE series_id = ?",
                       (series_id,)).fetchone()
    if not row:
        return None
    d = dict(row)
    d["report_rules"] = _rules(d.get("report_rules"))
    return d


def update_series(conn, series_id: str, **fields) -> dict:
    s = get_series(conn, series_id)
    if not s:
        raise ValueError(f"No series {series_id}.")
    allowed = {"name", "code", "organisation", "cadence_days", "language",
               "project", "notes", "status"}
    sets, vals = [], []
    for k, v in fields.items():
        if k not in allowed or v is None:
            continue
        if k == "code":
            v = str(v).strip().upper()
            if v:
                clash = conn.execute(
                    "SELECT name FROM meeting_series WHERE UPPER(code) = ? "
                    "AND series_id != ?", (v, series_id)).fetchone()
                if clash:
                    raise ValueError(f"Code {v} already belongs to '{clash['name']}'.")
        sets.append(f"{k} = ?")
        vals.append(v)
    if sets:
        sets.append("updated_at = ?")
        vals.extend([_now(), series_id])
        conn.execute(f"UPDATE meeting_series SET {', '.join(sets)} "
                     "WHERE series_id = ?", vals)
        conn.commit()
    return get_series(conn, series_id)


def add_report_rule(conn, series_id: str, rule: str) -> list[str]:
    """Append a standing report rule to ONE series. Duplicates are ignored."""
    s = get_series(conn, series_id)
    if not s:
        raise ValueError(f"No series {series_id}.")
    rule = re.sub(r"\s+", " ", rule or "").strip()
    rules = s["report_rules"]
    if rule and _fold(rule) not in {_fold(r) for r in rules}:
        rules.append(rule)
        conn.execute("UPDATE meeting_series SET report_rules = ?, updated_at = ? "
                     "WHERE series_id = ?",
                     (json.dumps(rules, ensure_ascii=False), _now(), series_id))
        conn.commit()
    return rules


def remove_report_rule(conn, series_id: str, index: int) -> list[str]:
    s = get_series(conn, series_id)
    if not s:
        raise ValueError(f"No series {series_id}.")
    rules = s["report_rules"]
    if 0 <= index < len(rules):
        rules.pop(index)
        conn.execute("UPDATE meeting_series SET report_rules = ?, updated_at = ? "
                     "WHERE series_id = ?",
                     (json.dumps(rules, ensure_ascii=False), _now(), series_id))
        conn.commit()
    return rules


def series_meetings(conn, series_id: str) -> list[dict]:
    ensure_schema(conn)
    rows = conn.execute(
        "SELECT meeting_id, title, meeting_date, sequence_label, location, "
        "attendees, report_path, previous_meeting_id FROM meetings "
        "WHERE series_id = ? ORDER BY meeting_date, created_at", (series_id,)
    ).fetchall()
    return [dict(r) for r in rows]


def list_series(conn, include_inactive: bool = False) -> list[dict]:
    """Every series with its last meeting, next expected date and open actions."""
    ensure_schema(conn)
    q = "SELECT * FROM meeting_series"
    if not include_inactive:
        q += " WHERE status = 'active'"
    out = []
    for r in conn.execute(q + " ORDER BY name").fetchall():
        d = dict(r)
        d["report_rules"] = _rules(d.get("report_rules"))
        ms = series_meetings(conn, d["series_id"])
        d["n_meetings"] = len(ms)
        d["last_meeting"] = ms[-1] if ms else None
        d["next_expected"] = _next_expected(d, ms)
        d["n_open_actions"] = len(open_actions(conn, d["series_id"]))
        out.append(d)
    return out


def _parse_date(s) -> _dt.date | None:
    try:
        return _dt.date.fromisoformat(str(s)[:10])
    except (TypeError, ValueError):
        return None


def _next_expected(series: dict, meetings: list[dict]) -> str:
    if not meetings or not series.get("cadence_days"):
        return ""
    last = _parse_date(meetings[-1].get("meeting_date"))
    if not last:
        return ""
    return (last + _dt.timedelta(days=int(series["cadence_days"]))).isoformat()


# ---------------------------------------------------------------------------
# Recognition
# ---------------------------------------------------------------------------

def _code_tokens(text: str) -> set[str]:
    """Upper-cased word tokens of a string, for an exact code comparison.

    Splits on anything that is not a letter or digit, so 'ABCD - Réunion' and
    '2026_T3_ABCD.pptx' both yield ABCD.
    """
    return {t.upper() for t in re.findall(r"[A-Za-z0-9]+", _fold(text or ""))}


_TITLE_NOISE = re.compile(
    r"\b(?:t[1-4]|q[1-4]|s[12]|20\d\d|19\d\d|\d+)\b")


def _title_tokens(title: str) -> set[str]:
    t = _TITLE_NOISE.sub(" ", _fold(title or ""))
    return {w for w in re.findall(r"[a-z]{3,}", t)}


def _person_series_counts(conn) -> tuple[dict[str, int], dict[str, set[str]]]:
    """How many series each person has attended, and each series' people."""
    rows = conn.execute(
        "SELECT m.series_id, a.person_id FROM meeting_attendance a "
        "JOIN meetings m ON m.meeting_id = a.meeting_id "
        "WHERE m.series_id IS NOT NULL AND m.series_id != ''").fetchall()
    people: dict[str, set[str]] = {}
    for r in rows:
        people.setdefault(r["series_id"], set()).add(r["person_id"])
    counts: dict[str, int] = {}
    for members in people.values():
        for p in members:
            counts[p] = counts.get(p, 0) + 1
    return counts, people


def match_series(conn, title: str = "", meeting_date: str = "",
                 attendees=None, filenames=None, text: str = "") -> dict:
    """Rank the existing series this meeting could continue.

    Returns {decision, candidates, question}. `decision` is:
      propose    one series is clearly ahead — ask the researcher to confirm it
      ambiguous  two or more are close — ask the researcher to choose
      new        nothing fits — offer to start a series

    Never links anything. Every candidate carries the reasons behind its score,
    so the question put to the researcher can say WHY.
    """
    ensure_schema(conn)
    series = [s for s in list_series(conn)]
    if not series:
        return {"decision": "new", "candidates": [],
                "question": "No meeting series exists yet. Start one for this meeting?"}

    head = " ".join([title or ""] + list(filenames or []))
    head_tokens = _code_tokens(head)
    text_tokens = _code_tokens((text or "")[:20000])
    people = {person_key(p) for p in parse_people(attendees)}
    people.discard("")
    counts, series_people = _person_series_counts(conn)
    when = _parse_date(meeting_date)
    t_tokens = _title_tokens(title)

    cands = []
    for s in series:
        score, reasons = 0.0, []
        code = (s.get("code") or "").upper()
        if code and code in head_tokens:
            score += _W_CODE_TITLE
            reasons.append(f"code {code} in the title or file names")
        elif code and code in text_tokens:
            score += _W_CODE_TEXT
            reasons.append(f"code {code} mentioned in the text")

        members = series_people.get(s["series_id"], set())
        if people and members:
            # A person seen in n series weighs 1/n: the researcher and colleagues
            # who attend every series carry almost nothing. Unknown people weigh
            # 1, so a room full of strangers lowers the fraction.
            w = {p: 1.0 / counts.get(p, 1) for p in people}
            total = sum(w.values())
            shared = [p for p in people if p in members]
            frac = sum(w[p] for p in shared) / total if total else 0.0
            if frac > 0:
                score += _W_PEOPLE * frac
                distinctive = [p for p in shared if counts.get(p, 1) == 1]
                reasons.append(
                    f"{len(shared)} participant(s) in common"
                    + (f", {len(distinctive)} seen only in this series"
                       if distinctive else ", none specific to this series"))

        last = _parse_date((s.get("last_meeting") or {}).get("meeting_date"))
        cad = int(s.get("cadence_days") or 0)
        if when and last and cad:
            gap = (when - last).days
            if gap > 0:
                fit = max(0.0, 1.0 - abs(gap - cad) / cad)
                if fit > 0:
                    score += _W_TIMING * fit
                    reasons.append(f"{gap} days after the last meeting "
                                   f"(cadence {cad})")

        st = _title_tokens(s.get("name", "")) | _title_tokens(
            (s.get("last_meeting") or {}).get("title", ""))
        if t_tokens and st:
            j = len(t_tokens & st) / len(t_tokens | st)
            if j > 0:
                score += _W_TITLE * j

        if score >= CANDIDATE_AT:
            lm = s.get("last_meeting") or {}
            cands.append({
                "series_id": s["series_id"], "name": s["name"],
                "code": s.get("code", ""), "score": round(score, 3),
                "reasons": reasons, "n_open_actions": s["n_open_actions"],
                "last_meeting_date": lm.get("meeting_date", ""),
                "last_meeting_label": lm.get("sequence_label", ""),
            })

    cands.sort(key=lambda c: c["score"], reverse=True)
    if not cands:
        return {"decision": "new", "candidates": [],
                "question": "This does not look like a continuation of any "
                            "existing series. Start a new series for it?"}
    top = cands[0]
    second = cands[1]["score"] if len(cands) > 1 else 0.0
    if top["score"] >= PROPOSE_AT and top["score"] - second >= MARGIN:
        last = top["last_meeting_label"] or top["last_meeting_date"] or "the last meeting"
        q = (f"This looks like the follow-up to {top['name']} ({last}), "
             f"with {top['n_open_actions']} open action(s) — "
             f"{'; '.join(top['reasons'])}. Correct?")
        return {"decision": "propose", "candidates": cands, "question": q}
    names = ", ".join(c["name"] for c in cands[:3])
    return {"decision": "ambiguous", "candidates": cands,
            "question": f"Which series does this meeting continue: {names}, "
                        "or a new one?"}


# ---------------------------------------------------------------------------
# Saving a meeting
# ---------------------------------------------------------------------------

def _record_attendance(conn, meeting_id: str, meeting_date: str,
                       people: list[str]) -> None:
    for name in people:
        pid = person_key(name)
        if not pid:
            continue
        conn.execute("INSERT OR IGNORE INTO meeting_attendance (meeting_id, person_id) "
                     "VALUES (?, ?)", (meeting_id, pid))
        row = conn.execute("SELECT meeting_count, last_meeting_date, name FROM "
                           "meeting_persons WHERE person_id = ?", (pid,)).fetchone()
        if row:
            last = max(str(row["last_meeting_date"] or ""), meeting_date or "")
            # Keep the longest spelling seen: transcripts shorten names more often
            # than they lengthen them.
            keep = name if len(name) > len(row["name"] or "") else row["name"]
            conn.execute("UPDATE meeting_persons SET meeting_count = meeting_count + 1, "
                         "last_meeting_date = ?, name = ? WHERE person_id = ?",
                         (last, keep, pid))
        else:
            conn.execute("INSERT INTO meeting_persons (person_id, name, "
                         "last_meeting_date, meeting_count) VALUES (?, ?, ?, 1)",
                         (pid, name, meeting_date))


def _norm_action(a) -> dict | None:
    if isinstance(a, str):
        a = {"description": a}
    if not isinstance(a, dict):
        return None
    desc = re.sub(r"\s+", " ", str(a.get("description") or a.get("action")
                                   or a.get("text") or "")).strip()
    if not desc:
        return None
    status = str(a.get("status") or "open").strip().lower()
    return {"description": desc[:1000],
            "owner": str(a.get("owner") or "").strip()[:300],
            "due_date": str(a.get("due_date") or a.get("due") or "").strip()[:60],
            "status": status if status in ACTION_STATUSES else "open"}


def previous_meeting(conn, series_id: str, before_date: str,
                     exclude: str = "") -> dict | None:
    row = conn.execute(
        "SELECT meeting_id, title, meeting_date, sequence_label FROM meetings "
        "WHERE series_id = ? AND meeting_date <= ? AND meeting_id != ? "
        "ORDER BY meeting_date DESC, created_at DESC LIMIT 1",
        (series_id, before_date or "9999", exclude)).fetchone()
    return dict(row) if row else None


def save_meeting(conn, title: str, meeting_date: str, attendees=None,
                 location: str = "", series_id: str = "", sequence_label: str = "",
                 summary: str = "", discussion=None, decisions=None, actions=None,
                 next_meeting: str = "", transcript: str = "",
                 attachments=None, language: str = "", meeting_type: str = "meeting",
                 project: str = "", meeting_id: str = "", follow_ups=None,
                 duration_minutes: int | None = None) -> str:
    """Insert one meeting with its participants and actions. Returns the id.

    `series_id` must already be confirmed by the researcher — this function links
    what it is told to link and never guesses.
    """
    ensure_schema(conn)
    title = (title or "").strip()
    if not title:
        raise ValueError("A meeting needs a title.")
    if not _parse_date(meeting_date):
        raise ValueError("meeting_date must be an ISO date (YYYY-MM-DD).")
    if series_id and not get_series(conn, series_id):
        raise ValueError(f"No series {series_id}.")
    people = parse_people(attendees)
    mid = meeting_id or f"mtg-{uuid.uuid4().hex[:12]}"
    now = _now()
    prev = previous_meeting(conn, series_id, meeting_date) if series_id else None
    acts = [x for x in (_norm_action(a) for a in (actions or [])) if x]
    disc = [str(d).strip() for d in (discussion or []) if str(d).strip()]
    decs = [str(d).strip() for d in (decisions or []) if str(d).strip()]
    fups = [str(d).strip() for d in (follow_ups or []) if str(d).strip()]

    conn.execute(
        "INSERT INTO meetings (meeting_id, title, meeting_date, attendees, "
        "meeting_type, project, transcript, transcript_status, action_items, "
        "decisions, follow_ups, series_id, sequence_label, previous_meeting_id, "
        "location, summary, discussion, attachments, language, next_meeting, "
        "linked_meetings, duration_minutes, status, created_at) "
        "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?, 'filed', ?)",
        # Attendees stay a comma-separated string: that is what every existing
        # dashboard query and template reads, and parse_people reads both.
        (mid, title, meeting_date[:10], ", ".join(people),
         meeting_type, project, transcript or "",
         "filed" if transcript else "none",
         json.dumps([a["description"] for a in acts], ensure_ascii=False),
         json.dumps(decs, ensure_ascii=False), json.dumps(fups, ensure_ascii=False),
         series_id or "",
         sequence_label.strip(), (prev or {}).get("meeting_id", ""),
         location.strip(), summary.strip(), json.dumps(disc, ensure_ascii=False),
         json.dumps(list(attachments or []), ensure_ascii=False), language.strip(),
         next_meeting.strip(),
         json.dumps([prev["meeting_id"]] if prev else []), duration_minutes, now))

    for a in acts:
        conn.execute(
            "INSERT INTO meeting_actions (meeting_id, description, status, due_date, "
            "created_at, owner, series_id, updated_at) VALUES (?,?,?,?,?,?,?,?)",
            (mid, a["description"], a["status"], a["due_date"], now, a["owner"],
             series_id or "", now))
    _record_attendance(conn, mid, meeting_date[:10], people)
    conn.commit()
    return mid


def link_meeting_to_series(conn, meeting_id: str, series_id: str,
                           sequence_label: str = "") -> dict:
    """Attach an already-saved meeting to a series, after the researcher confirmed."""
    ensure_schema(conn)
    m = conn.execute("SELECT meeting_id, meeting_date FROM meetings WHERE meeting_id = ?",
                     (meeting_id,)).fetchone()
    if not m:
        raise ValueError(f"No meeting {meeting_id}.")
    if not get_series(conn, series_id):
        raise ValueError(f"No series {series_id}.")
    prev = previous_meeting(conn, series_id, m["meeting_date"], exclude=meeting_id)
    conn.execute(
        "UPDATE meetings SET series_id = ?, previous_meeting_id = ?, "
        "linked_meetings = ?, sequence_label = CASE WHEN ? != '' THEN ? "
        "ELSE sequence_label END WHERE meeting_id = ?",
        (series_id, (prev or {}).get("meeting_id", ""),
         json.dumps([prev["meeting_id"]] if prev else []),
         sequence_label, sequence_label, meeting_id))
    conn.execute("UPDATE meeting_actions SET series_id = ? WHERE meeting_id = ?",
                 (series_id, meeting_id))
    conn.commit()
    return {"meeting_id": meeting_id, "series_id": series_id,
            "previous_meeting_id": (prev or {}).get("meeting_id", "")}


def add_attachment(conn, meeting_id: str, path: str) -> list[str]:
    ensure_schema(conn)
    row = conn.execute("SELECT attachments FROM meetings WHERE meeting_id = ?",
                       (meeting_id,)).fetchone()
    if not row:
        raise ValueError(f"No meeting {meeting_id}.")
    try:
        items = json.loads(row["attachments"] or "[]")
    except ValueError:
        items = []
    if path not in items:
        items.append(path)
        conn.execute("UPDATE meetings SET attachments = ? WHERE meeting_id = ?",
                     (json.dumps(items, ensure_ascii=False), meeting_id))
        conn.commit()
    return items


# ---------------------------------------------------------------------------
# Actions across meetings
# ---------------------------------------------------------------------------

def _last_review(conn, action_id: int) -> dict | None:
    r = conn.execute(
        "SELECT r.status, r.note, r.meeting_id, r.reviewed_at, m.sequence_label, "
        "m.meeting_date FROM meeting_action_reviews r "
        "LEFT JOIN meetings m ON m.meeting_id = r.meeting_id "
        "WHERE r.action_id = ? ORDER BY r.reviewed_at DESC, r.id DESC LIMIT 1",
        (action_id,)).fetchone()
    return dict(r) if r else None


def series_actions(conn, series_id: str, include_closed: bool = True) -> list[dict]:
    """Every action raised in a series, oldest meeting first, with its history."""
    ensure_schema(conn)
    rows = conn.execute(
        "SELECT a.id, a.meeting_id, a.description, a.owner, a.due_date, a.status, "
        "a.status_note, m.meeting_date, m.sequence_label, m.title AS meeting_title "
        "FROM meeting_actions a JOIN meetings m ON m.meeting_id = a.meeting_id "
        "WHERE m.series_id = ? ORDER BY m.meeting_date, a.id", (series_id,)
    ).fetchall()
    out = []
    for r in rows:
        d = dict(r)
        if not include_closed and (d["status"] or "open") in _CLOSED:
            continue
        d["last_review"] = _last_review(conn, d["id"])
        out.append(d)
    return out


def open_actions(conn, series_id: str) -> list[dict]:
    return series_actions(conn, series_id, include_closed=False)


def review_action(conn, action_id: int, status: str, meeting_id: str = "",
                  note: str = "") -> dict:
    """Record an action's status, optionally AT a given meeting.

    The action keeps one identity for its whole life. Reviewing it at the next
    meeting adds a review row and updates the action's current status; nothing
    is copied, so "how long has this been open" stays answerable.
    """
    ensure_schema(conn)
    status = (status or "").strip().lower().replace(" ", "_")
    if status not in ACTION_STATUSES:
        raise ValueError(f"status must be one of {', '.join(ACTION_STATUSES)}")
    a = conn.execute("SELECT id FROM meeting_actions WHERE id = ?",
                     (int(action_id),)).fetchone()
    if not a:
        raise ValueError(f"No action {action_id}.")
    now = _now()
    if meeting_id:
        if not conn.execute("SELECT 1 FROM meetings WHERE meeting_id = ?",
                            (meeting_id,)).fetchone():
            raise ValueError(f"No meeting {meeting_id}.")
        conn.execute("INSERT INTO meeting_action_reviews (action_id, meeting_id, "
                     "status, note, reviewed_at) VALUES (?,?,?,?,?)",
                     (int(action_id), meeting_id, status, note.strip(), now))
    conn.execute("UPDATE meeting_actions SET status = ?, status_note = ?, "
                 "updated_at = ? WHERE id = ?",
                 (status, note.strip(), now, int(action_id)))
    conn.commit()
    return {"action_id": int(action_id), "status": status, "meeting_id": meeting_id}


# ---------------------------------------------------------------------------
# Briefing and report
# ---------------------------------------------------------------------------

_LABELS = {
    "fr": {
        "report": "Compte rendu", "date": "Date", "place": "Lieu",
        "participants": "Participants", "summary": "Synthèse",
        "previous": "Suivi des actions de la réunion précédente",
        "discussion": "Points de discussion", "decisions": "Décisions",
        "actions": "Plan d'actions", "num": "N°", "action": "Action",
        "owner": "Responsable", "due": "Échéance", "status": "Statut",
        "next": "Prochaine réunion", "none": "Aucune",
        "status_values": {"open": "Ouvert", "in_progress": "En cours",
                          "done": "Réalisé", "dropped": "Abandonné"},
        "raised": "Issue de",
    },
    "en": {
        "report": "Meeting report", "date": "Date", "place": "Place",
        "participants": "Participants", "summary": "Summary",
        "previous": "Follow-up of the previous meeting's actions",
        "discussion": "Discussion points", "decisions": "Decisions",
        "actions": "Action plan", "num": "No.", "action": "Action",
        "owner": "Owner", "due": "Due", "status": "Status",
        "next": "Next meeting", "none": "None",
        "status_values": {"open": "Open", "in_progress": "In progress",
                          "done": "Done", "dropped": "Dropped"},
        "raised": "Raised at",
    },
}


def labels(language: str) -> dict:
    return _LABELS.get((language or "en")[:2].lower(), _LABELS["en"])


def _loads(raw, default):
    try:
        v = json.loads(raw) if raw else default
        return v if isinstance(v, type(default)) else default
    except (TypeError, ValueError):
        return default


def get_meeting(conn, meeting_id: str) -> dict | None:
    ensure_schema(conn)
    r = conn.execute("SELECT * FROM meetings WHERE meeting_id = ?",
                     (meeting_id,)).fetchone()
    if not r:
        return None
    d = dict(r)
    d["attendees_list"] = parse_people(d.get("attendees"))
    d["discussion_list"] = _loads(d.get("discussion"), [])
    d["decisions_list"] = [x for x in _loads(d.get("decisions"), [])
                           if isinstance(x, str)]
    d["attachments_list"] = _loads(d.get("attachments"), [])
    d["actions"] = [dict(a) for a in conn.execute(
        "SELECT id, description, owner, due_date, status, status_note "
        "FROM meeting_actions WHERE meeting_id = ? ORDER BY id", (meeting_id,))]
    return d


def briefing(conn, series_id: str) -> str:
    """What to have in front of you before the next meeting of a series."""
    s = get_series(conn, series_id)
    if not s:
        raise ValueError(f"No series {series_id}.")
    ms = series_meetings(conn, series_id)
    lines = [f"# Briefing — {s['name']}", ""]
    if s.get("code"):
        lines.append(f"Code: {s['code']}")
    if s.get("cadence_days"):
        nxt = _next_expected(s, ms)
        lines.append(f"Cadence: every {s['cadence_days']} days"
                     + (f" · next expected around {nxt}" if nxt else ""))
    lines.append(f"Meetings held: {len(ms)}")
    if not ms:
        lines += ["", "No meeting has been recorded in this series yet."]
        return "\n".join(lines)
    last = get_meeting(conn, ms[-1]["meeting_id"])
    lines += ["", f"## Last meeting — {last.get('sequence_label') or last['title']} "
                  f"({last['meeting_date']})"]
    if last.get("location"):
        lines.append(f"Place: {last['location']}")
    if last["attendees_list"]:
        lines.append("Participants: " + ", ".join(last["attendees_list"]))
    if last.get("summary"):
        lines += ["", last["summary"]]
    if last["decisions_list"]:
        lines += ["", "### Decisions"] + [f"- {d}" for d in last["decisions_list"]]
    if last["attachments_list"]:
        lines += ["", "### Files from that meeting (compare the figures)"] + [
            f"- {p}" for p in last["attachments_list"]]

    acts = open_actions(conn, series_id)
    today = _dt.date.today()
    lines += ["", f"## Open actions ({len(acts)})"]
    if acts:
        lines += ["| # | Action | Owner | Due | Raised | Last status |",
                  "|---|---|---|---|---|---|"]
        for a in acts:
            due = a.get("due_date") or ""
            d = _parse_date(due)
            flag = " ⚠ overdue" if d and d < today else ""
            lr = a.get("last_review")
            last_status = (f"{lr['status']} at {lr.get('sequence_label') or lr.get('meeting_date') or ''}"
                           if lr else a.get("status") or "open")
            lines.append(
                f"| {a['id']} | {a['description']} | {a.get('owner') or '—'} | "
                f"{due or '—'}{flag} | {a.get('sequence_label') or a['meeting_date']} | "
                f"{last_status} |")
    else:
        lines.append("No open actions.")
    if s["report_rules"]:
        lines += ["", "## Report rules for this series"] + [
            f"- {r}" for r in s["report_rules"]]
    return "\n".join(lines)


def draft_report(conn, meeting_id: str) -> str:
    """Assemble the report from what is stored, in the meeting's language.

    Deterministic on purpose: the model writes the prose of the summary and the
    discussion points when it saves the meeting; this lays them out, adds the
    previous meeting's actions with their status AT this meeting, and the new
    action plan. The same input always gives the same document.
    """
    m = get_meeting(conn, meeting_id)
    if not m:
        raise ValueError(f"No meeting {meeting_id}.")
    s = get_series(conn, m["series_id"]) if m.get("series_id") else None
    lang = m.get("language") or (s or {}).get("language") or "en"
    L = labels(lang)
    sv = L["status_values"]
    title = m["title"]
    # French typography puts a space before the colon.
    sep = " : " if lang.startswith("fr") else ": "
    out = [f"# {L['report']} — {title}", "",
           f"**{L['date']}**{sep}{m['meeting_date']}"]
    if m.get("location"):
        out.append(f"**{L['place']}**{sep}{m['location']}")
    if m["attendees_list"]:
        out.append(f"**{L['participants']}**{sep}" + ", ".join(m["attendees_list"]))
    if m.get("summary"):
        out += ["", f"## {L['summary']}", "", m["summary"]]

    prev_rows = []
    if s and m.get("previous_meeting_id"):
        for a in series_actions(conn, s["series_id"]):
            if a["meeting_id"] == meeting_id:
                continue
            if (a["meeting_date"] or "") > m["meeting_date"]:
                continue
            here = conn.execute(
                "SELECT status, note FROM meeting_action_reviews WHERE action_id = ? "
                "AND meeting_id = ? ORDER BY id DESC LIMIT 1",
                (a["id"], meeting_id)).fetchone()
            if here:
                prev_rows.append((a, here["status"], here["note"] or ""))
            elif (a["status"] or "open") not in _CLOSED:
                prev_rows.append((a, a["status"] or "open", a.get("status_note") or ""))
    if prev_rows:
        out += ["", f"## {L['previous']}", "",
                f"| {L['num']} | {L['action']} | {L['owner']} | {L['raised']} | {L['status']} |",
                "|---|---|---|---|---|"]
        for i, (a, st, note) in enumerate(prev_rows, 1):
            stxt = sv.get(st, st) + (f" — {note}" if note else "")
            out.append(f"| {i} | {a['description']} | {a.get('owner') or '—'} | "
                       f"{a.get('sequence_label') or a['meeting_date']} | {stxt} |")

    if m["discussion_list"]:
        out += ["", f"## {L['discussion']}", ""] + [f"- {d}" for d in m["discussion_list"]]
    if m["decisions_list"]:
        out += ["", f"## {L['decisions']}", ""] + [f"- {d}" for d in m["decisions_list"]]
    out += ["", f"## {L['actions']}", ""]
    if m["actions"]:
        out += [f"| {L['num']} | {L['action']} | {L['owner']} | {L['due']} | {L['status']} |",
                "|---|---|---|---|---|"]
        for i, a in enumerate(m["actions"], 1):
            out.append(f"| {i} | {a['description']} | {a.get('owner') or '—'} | "
                       f"{a.get('due_date') or '—'} | {sv.get(a['status'] or 'open', a['status'])} |")
    else:
        out.append(L["none"])
    if m.get("next_meeting"):
        out += ["", f"## {L['next']}", "", m["next_meeting"]]
    return "\n".join(out) + "\n"


def report_diff(draft: str, final: str) -> dict:
    """What the researcher removed, added and rewrote between draft and final.

    No model here. The lines are handed to the conversation, which turns them
    into report rules (`add_report_rule`) — a rule is a judgement, and a diff is
    evidence for it, not the rule itself.
    """
    def _plain(line: str) -> str:
        # The final usually comes back from Word, without the draft's Markdown.
        # Compare words, not markup, or every line reads as rewritten.
        line = line.strip()
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            line = " ".join(c for c in cells if c)
        line = re.sub(r"^(#{1,6}\s+|[-*•]\s+)", "", line)
        line = line.replace("**", "")
        return re.sub(r"\s+", " ", line).strip()

    def _lines(t):
        out = []
        for x in (t or "").splitlines():
            if re.fullmatch(r"\s*\|?[\s:|-]*-{2,}[\s:|-]*\|?\s*", x):
                continue          # a table's separator row
            p = _plain(x)
            if p:
                out.append(p)
        return out
    a, b = _lines(draft), _lines(final)
    removed, added, changed = [], [], []
    sm = difflib.SequenceMatcher(a=a, b=b, autojunk=False)
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == "delete":
            removed += a[i1:i2]
        elif op == "insert":
            added += b[j1:j2]
        elif op == "replace":
            # Pair each new line with the most similar old line in the block,
            # not with the one at the same position: a removed line before a
            # rewritten one would otherwise shift every pairing by one.
            old, new = list(a[i1:i2]), b[j1:j2]
            for n in new:
                best, best_r = None, 0.0
                for o in old:
                    r = 1.0 if (n in o or o in n) else \
                        difflib.SequenceMatcher(a=o, b=n).ratio()
                    if r > best_r:
                        best, best_r = o, r
                if best is not None and best_r >= 0.5:
                    changed.append({"before": best, "after": n})
                    old.remove(best)
                else:
                    added.append(n)
            removed += old
    return {"removed": removed, "added": added, "changed": changed,
            "similarity": round(sm.ratio(), 3)}


# ---------------------------------------------------------------------------
# Word export
# ---------------------------------------------------------------------------

def _md_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for ln in lines:
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            continue
        rows.append(cells)
    return rows


def _add_runs(par, text: str) -> None:
    """Honour **bold** in a line; everything else is plain text."""
    for i, part in enumerate(re.split(r"\*\*", text)):
        if part:
            run = par.add_run(part)
            run.bold = bool(i % 2)


def markdown_to_docx(markdown: str, target_path: str) -> str:
    """Render a report written in the subset of Markdown `draft_report` produces.

    Compact on purpose — narrow margins and a 10 pt body — because a report of
    this kind is expected to fit on two pages and Word's defaults spread it to
    four. Tables become real Word tables, which the generic document tool did not
    do: it wrote a Markdown table into the file as lines of pipes.
    """
    from docx import Document
    from docx.shared import Cm, Pt

    doc = Document()
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(1.6)
        sec.left_margin = sec.right_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(2)

    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        raw = lines[i].rstrip()
        s = raw.strip()
        if not s:
            i += 1
            continue
        if s.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            rows = _md_table(block)
            if rows:
                ncol = max(len(r) for r in rows)
                t = doc.add_table(rows=len(rows), cols=ncol)
                t.style = "Table Grid"
                t.autofit = False
                # Equal columns squeeze the one column that carries the text.
                # Weight it: a number column narrow, the action wide.
                head = [(rows[0][c] if c < len(rows[0]) else "").lower()
                        for c in range(ncol)]
                weights = [0.6 if len(h) <= 3 else 5.0 if "action" in h else 2.0
                           for h in head]
                usable = 21.0 - 3.6       # A4 width minus the margins above, in cm
                widths = [Cm(usable * w / sum(weights)) for w in weights]
                # Word reads the cell widths; LibreOffice reads the grid.
                for ci, col in enumerate(t.columns):
                    col.width = widths[ci]
                for ri, r in enumerate(rows):
                    for ci in range(ncol):
                        cell = t.cell(ri, ci)
                        cell.width = widths[ci]
                        cell.text = ""
                        p = cell.paragraphs[0]
                        run = p.add_run(r[ci] if ci < len(r) else "")
                        run.font.size = Pt(9)
                        run.bold = ri == 0
            continue
        if s.startswith("# "):
            h = doc.add_heading(s[2:], level=1)
        elif s.startswith("## "):
            h = doc.add_heading(s[3:], level=2)
        elif s.startswith("### "):
            h = doc.add_heading(s[4:], level=3)
        elif s.startswith(("- ", "* ", "• ")):
            _add_runs(doc.add_paragraph(style="List Bullet"), s[2:])
            h = None
        else:
            _add_runs(doc.add_paragraph(), s)
            h = None
        if h is not None:
            h.paragraph_format.space_before = Pt(6)
            h.paragraph_format.space_after = Pt(2)
        i += 1
    doc.save(target_path)
    return target_path
